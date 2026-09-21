"""
CEKAS Excel Financial Report Generator
Generates multi-sheet Excel workbook with official SMA Kartika XIX-1 Bandung formatting.
100% offline, pure Python with openpyxl.
"""

from io import BytesIO
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .database import CekasDB
from .formatters import format_rupiah, format_wib
from .config import CLASS_ID, CLASS_NAME, SCHOOL_NAME, SCHOOL_YEAR, POCKET_CONFIG

def generate_excel_report(db: CekasDB, class_id: str = CLASS_ID) -> BytesIO:
    wb = openpyxl.Workbook()
    
    # ----------------------------------------------------
    # Styles Setup (Navy & Pastel theme)
    # ----------------------------------------------------
    navy_header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=16, bold=True, color="0F172A")
    subtitle_font = Font(name="Segoe UI", size=11, italic=True, color="475569")
    bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    regular_font = Font(name="Segoe UI", size=10, color="1E293B")
    
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    in_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Mint
    out_fill = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid") # Coral
    corr_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid") # Lemon
    
    # ====================================================
    # SHEET 1: RINGKASAN SALDO & POS ANGGARAN
    # ====================================================
    ws1 = wb.active
    ws1.title = "Ringkasan Kas"
    ws1.views.sheetView[0].showGridLines = True
    
    cls_info = db.get_or_create_class(class_id)
    saldo_total = cls_info["saldo"]
    alokasi = cls_info["alokasi"]
    
    ws1["A1"] = f"LAPORAN KEUANGAN KAS DIGITAL {CLASS_ID}"
    ws1["A1"].font = title_font
    ws1["A2"] = f"{SCHOOL_NAME} • TAHUN AJARAN {SCHOOL_YEAR}"
    ws1["A2"].font = subtitle_font
    ws1["A3"] = f"Dicetak pada: {format_wib()}"
    ws1["A3"].font = subtitle_font
    
    ws1["A5"] = "TOTAL SALDO KAS SAAT INI"
    ws1["A5"].font = bold_font
    ws1["B5"] = saldo_total
    ws1["B5"].font = Font(name="Segoe UI", size=14, bold=True, color="059669")
    ws1["B5"].number_format = '"Rp "#,##0'
    
    # Table Pos Anggaran
    headers_pos = ["Kode Pos", "Nama Pos Alokasi", "Nominal Saldo", "Persentase", "Fungsi & Peruntukan"]
    for col_idx, h in enumerate(headers_pos, start=1):
        cell = ws1.cell(row=7, column=col_idx, value=h)
        cell.fill = navy_header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border
        
    row_idx = 8
    for p_key, meta in POCKET_CONFIG.items():
        p_saldo = alokasi.get(p_key, 0)
        pct = (p_saldo / saldo_total) if saldo_total > 0 else 0
        
        c1 = ws1.cell(row=row_idx, column=1, value=p_key.upper())
        c2 = ws1.cell(row=row_idx, column=2, value=f"{meta['icon']} {meta['label']}")
        c3 = ws1.cell(row=row_idx, column=3, value=p_saldo)
        c4 = ws1.cell(row=row_idx, column=4, value=pct)
        c5 = ws1.cell(row=row_idx, column=5, value=meta["description"])
        
        c1.alignment = Alignment(horizontal="center")
        c3.number_format = '"Rp "#,##0'
        c4.number_format = '0.0%'
        
        for c in [c1, c2, c3, c4, c5]:
            c.font = regular_font
            c.border = cell_border
            
        row_idx += 1
        
    # Total row
    c_tot_label = ws1.cell(row=row_idx, column=2, value="TOTAL AKUMULASI")
    c_tot_label.font = bold_font
    c_tot_val = ws1.cell(row=row_idx, column=3, value=f"=SUM(C8:C{row_idx-1})")
    c_tot_val.font = bold_font
    c_tot_val.number_format = '"Rp "#,##0'
    c_tot_pct = ws1.cell(row=row_idx, column=4, value=f"=SUM(D8:D{row_idx-1})")
    c_tot_pct.font = bold_font
    c_tot_pct.number_format = '0.0%'
    
    for c in [ws1.cell(row=row_idx, column=1), c_tot_label, c_tot_val, c_tot_pct, ws1.cell(row=row_idx, column=5)]:
        c.border = cell_border
        
    # Auto adjust column widths
    for col in ws1.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 14)
        
    # ====================================================
    # SHEET 2: BUKU KAS UMUM (LEDGER TRANSAKSI)
    # ====================================================
    ws2 = wb.create_sheet(title="Buku Kas Umum")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2["A1"] = f"BUKU KAS UMUM (MUTASI LENGKAP) {CLASS_ID}"
    ws2["A1"].font = title_font
    ws2["A2"] = "Audit Trail Append-Only Transaksi Digital"
    ws2["A2"].font = subtitle_font
    
    headers_tx = ["ID", "Waktu", "Tipe", "Pos Anggaran", "Deskripsi / Alasan", "Masuk (Debit)", "Keluar (Kredit)", "Pencatat"]
    for col_idx, h in enumerate(headers_tx, start=1):
        cell = ws2.cell(row=4, column=col_idx, value=h)
        cell.fill = navy_header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border
        
    txs = db.get_recent_transactions(class_id, limit=500)
    txs.reverse() # chronological order for ledger
    
    row_tx = 5
    for t in txs:
        t_id = f"#{t['id']}"
        t_time = str(t["created_at"])[:19]
        t_type = "📥 MASUK" if t["type"] == "in" else ("📤 KELUAR" if t["type"] == "out" else "↩️ KOREKSI")
        t_pos = POCKET_CONFIG.get(t["category"], {}).get("label", t["category"])
        t_desc = t["description"]
        debit = t["amount"] if t["type"] == "in" else None
        kredit = t["amount"] if t["type"] in ["out", "correction"] else None
        pencatat = t["input_by"]
        
        c1 = ws2.cell(row=row_tx, column=1, value=t_id)
        c2 = ws2.cell(row=row_tx, column=2, value=t_time)
        c3 = ws2.cell(row=row_tx, column=3, value=t_type)
        c4 = ws2.cell(row=row_tx, column=4, value=t_pos)
        c5 = ws2.cell(row=row_tx, column=5, value=t_desc)
        c6 = ws2.cell(row=row_tx, column=6, value=debit)
        c7 = ws2.cell(row=row_tx, column=7, value=kredit)
        c8 = ws2.cell(row=row_tx, column=8, value=pencatat)
        
        c1.alignment = Alignment(horizontal="center")
        c2.alignment = Alignment(horizontal="center")
        c3.alignment = Alignment(horizontal="center")
        c6.number_format = '"Rp "#,##0'
        c7.number_format = '"Rp "#,##0'
        
        if t["type"] == "in":
            c3.fill = in_fill
        elif t["type"] == "out":
            c3.fill = out_fill
        else:
            c3.fill = corr_fill
            
        for c in [c1, c2, c3, c4, c5, c6, c7, c8]:
            c.font = regular_font
            c.border = cell_border
            
        row_tx += 1
        
    # Auto width sheet 2
    for col in ws2.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
